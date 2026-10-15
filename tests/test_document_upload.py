import io

import pytest
from docx import Document as DocxDocument

from backend.services.ai_service import analyze_ticket
from backend.services.routing_service import routing_logic
from backend.utils.document_parser import Block, chunk_blocks, parse_document

# A small printer guide: nothing in the 20 seed entries covers printers, so if a
# printer question retrieves this content it can only have come from the upload.
TITLE = "Kyocera TASKalfa 3253ci printer guide"
SECTIONS = [
    (
        "Clearing a paper jam",
        "Lift the green release lever and open the right side cover. Pull the jammed "
        "sheet out slowly in the direction the paper travels so it does not tear. "
        "Close the cover firmly and the printer resumes the job on its own.",
    ),
    (
        "Replacing the toner cartridge",
        "Open the front cover, turn the empty cartridge lock to the left and slide the "
        "cartridge out. Shake the new cartridge five times before inserting it.",
    ),
    (
        "Adding the printer to a laptop",
        "Open printer settings, choose Add printer by IP address and enter 10.20.4.15. "
        "Select the KX driver when asked and print a test page to confirm.",
    ),
]
JAM_MARKER = "green release lever"


def make_txt() -> bytes:
    parts = [f"# {TITLE}"]
    for heading, body in SECTIONS:
        parts += [f"## {heading}", body]
    return "\n\n".join(parts).encode("utf-8")


def make_docx() -> bytes:
    document = DocxDocument()
    document.add_heading(TITLE, level=0)
    for heading, body in SECTIONS:
        document.add_heading(heading, level=1)
        document.add_paragraph(body)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pdf() -> bytes:
    """Minimal one-page text PDF, written by hand to avoid a PDF-writer dependency."""
    lines = [TITLE, ""]
    for heading, body in SECTIONS:
        lines.append(heading)
        line = ""
        for word in body.split():
            if line and len(line) + 1 + len(word) > 85:
                lines.append(line)
                line = word
            else:
                line = f"{line} {word}" if line else word
        lines += [line, ""]

    def escape(text):
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    stream = "BT /F1 10 Tf 14 TL 50 780 Td " + " ".join(
        f"({escape(line)}) Tj T*" for line in lines
    ) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] "
        "/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
    ]
    out = "%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    out += "".join(f"{offset:010d} 00000 n \n" for offset in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n"
    return out.encode("latin-1")


SAMPLES = {
    "printer-guide.txt": make_txt,
    "printer-guide.docx": make_docx,
    "printer-guide.pdf": make_pdf,
}


def retrieved_context_parts(llm_call) -> list[str]:
    """The ranked context analyze_ticket put in the prompt, one entry per hit."""
    user_prompt = llm_call["messages"][1]["content"]
    context = user_prompt.split("Relevant past solutions:\n", 1)[1]
    context = context.split("\n\nWrite a pinpoint support reply", 1)[0]
    return context.split("\n---\n")


@pytest.mark.parametrize("filename", list(SAMPLES))
def test_uploaded_document_is_chunked_embedded_and_retrieved(
    filename, client, collection, captured_llm_calls
):
    response = client.post(
        "/knowledge-base/upload",
        files={"file": (filename, SAMPLES[filename](), "application/octet-stream")},
    )
    assert response.status_code == 200, response.text
    chunks_added = response.json()["chunks_added"]
    assert chunks_added >= 1

    # Chunks are stored in the shared collection with real embeddings.
    stored = collection.get(
        where={"filename": filename}, include=["documents", "metadatas", "embeddings"]
    )
    assert len(stored["ids"]) == chunks_added
    assert all(meta["source"] == "document" for meta in stored["metadatas"])
    assert all(meta["reply"] == doc for meta, doc in zip(stored["metadatas"], stored["documents"]))
    assert all(len(embedding) == 384 for embedding in stored["embeddings"])  # all-MiniLM-L6-v2
    # Headings survive parsing in every format, so each section is its own chunk.
    assert chunks_added == len(SECTIONS)
    assert any(doc.startswith("Clearing a paper jam\n") for doc in stored["documents"])

    # The unchanged RAG flow ranks the uploaded content first for a paraphrased
    # question and passes it to the LLM as context.
    analysis = analyze_ticket("Paper is stuck inside the Kyocera printer, how do I get it out?")
    assert len(captured_llm_calls) == 1
    context = retrieved_context_parts(captured_llm_calls[0])
    assert JAM_MARKER in context[0]
    assert 0.0 <= analysis["confidence"] <= 1.0


def test_unrelated_question_still_retrieves_seed_entry(client, captured_llm_calls):
    client.post(
        "/knowledge-base/upload",
        files={"file": ("printer-guide.txt", make_txt(), "text/plain")},
    )
    analyze_ticket("How do I reset my password?")
    context = retrieved_context_parts(captured_llm_calls[0])
    assert context[0].startswith("Go to Settings > Security > Reset Password")


def test_reupload_replaces_previous_chunks(client, collection):
    files = {"file": ("replace-me.txt", make_txt(), "text/plain")}
    client.post("/knowledge-base/upload", files=files)
    client.post("/knowledge-base/upload", files=files)
    assert len(collection.get(where={"filename": "replace-me.txt"})["ids"]) == len(SECTIONS)

    files = {"file": ("replace-me.txt", b"Just one short paragraph now.", "text/plain")}
    client.post("/knowledge-base/upload", files=files)
    stored = collection.get(where={"filename": "replace-me.txt"})
    assert stored["documents"] == ["Just one short paragraph now."]


@pytest.mark.parametrize(
    "filename, payload, status",
    [
        ("notes.md", b"# hi", 415),
        ("empty.txt", b"  \n\n  ", 422),
        ("broken.pdf", b"not a pdf", 422),
    ],
)
def test_upload_rejects_bad_input(client, filename, payload, status):
    response = client.post("/knowledge-base/upload", files={"file": (filename, payload, "text/plain")})
    assert response.status_code == status


def test_chunks_respect_max_size_and_section_boundaries():
    long_paragraph = " ".join(f"Sentence number {i} explains one step." for i in range(80))
    blocks = [
        Block("Section A", is_heading=True),
        Block("Short intro."),
        Block(long_paragraph),
        Block("Section B", is_heading=True),
        Block("Only paragraph in B."),
    ]
    chunks = chunk_blocks(blocks, max_chars=300)
    assert all(len(chunk) <= 300 for chunk in chunks)
    assert all(chunk.startswith("Section A\n") for chunk in chunks[:-1])
    assert chunks[-1] == "Section B\nOnly paragraph in B."
    # Oversized paragraphs split at sentence ends, not mid-word.
    assert all(chunk.rstrip().endswith(".") for chunk in chunks)


def test_txt_heading_directly_followed_by_body():
    text = b"# Guide\n## Clearing a jam\nLift the lever.\nPull the sheet out.\n## Toner\nShake it."
    blocks = parse_document("guide.txt", text)
    assert [(b.text, b.is_heading) for b in blocks] == [
        ("Guide", True),
        ("Clearing a jam", True),
        ("Lift the lever. Pull the sheet out.", False),
        ("Toner", True),
        ("Shake it.", False),
    ]
    assert chunk_blocks(blocks) == ["Clearing a jam\nLift the lever. Pull the sheet out.", "Toner\nShake it."]


def test_txt_plain_headings_without_markdown():
    text = (
        b"Dock guide\n\n"
        b"External monitors not detected\nUnplug the cable and plug it back in.\n\n"
        b"Guest Wi-Fi\n\nVisitors use Corp-Guest.\nCodes last 24 hours.\n\n"
        b"Thanks"
    )
    assert chunk_blocks(parse_document("dock.txt", text)) == [
        "External monitors not detected\nUnplug the cable and plug it back in.",
        "Guest Wi-Fi\nVisitors use Corp-Guest. Codes last 24 hours.\nThanks",
    ]


def test_pdf_parser_joins_wrapped_lines():
    text = " ".join(block.text for block in parse_document("guide.pdf", make_pdf()))
    assert "Pull the jammed sheet out slowly in the direction the paper travels" in text


def _cosine_confidence(collection, question):
    """The original formula: 0.75 x top similarity + 0.25 x mean of the top 3."""
    distances = collection.query(query_texts=[question], n_results=3)["distances"][0]
    similarities = [max(0, 1 - d) for d in distances]
    return 0.75 * similarities[0] + 0.25 * sum(similarities) / len(similarities)


def test_question_answered_by_document_reaches_tier1(client, collection, captured_llm_calls):
    client.post(
        "/knowledge-base/upload",
        files={"file": ("printer-guide.txt", make_txt(), "text/plain")},
    )
    question = "How do I replace the toner on the office printer?"
    # Similarity alone leaves a passage-backed answer short of auto-resolve...
    assert routing_logic(_cosine_confidence(collection, question), 0.0)[0] != "tier1"
    # ...the answer check on the retrieved chunk lifts it.
    confidence = analyze_ticket(question)["confidence"]
    assert routing_logic(confidence, 0.0)[0] == "tier1"
    # A very negative customer still gets a human look, as before.
    assert routing_logic(confidence, -0.9)[0] != "tier1"


@pytest.mark.parametrize(
    "question",
    [
        "The printer is making a loud grinding noise",
        "Printer shows error code C6000 and won't start",
        "Can I print in colour on A3 paper?",
    ],
)
def test_same_topic_but_unanswered_is_not_auto_resolved(client, captured_llm_calls, question):
    client.post(
        "/knowledge-base/upload",
        files={"file": ("printer-guide.txt", make_txt(), "text/plain")},
    )
    assert routing_logic(analyze_ticket(question)["confidence"], 0.0)[0] != "tier1"


@pytest.mark.parametrize(
    "question",
    ["How do I reset my password?", "VPN won't connect on my MacBook", "my emails are not arriving"],
)
def test_seed_confidence_unchanged(client, collection, captured_llm_calls, question):
    client.post(
        "/knowledge-base/upload",
        files={"file": ("printer-guide.txt", make_txt(), "text/plain")},
    )
    expected = round(_cosine_confidence(collection, question), 4)
    assert analyze_ticket(question)["confidence"] == pytest.approx(expected, abs=1e-4)
